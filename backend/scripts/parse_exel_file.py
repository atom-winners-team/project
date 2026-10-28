#!/usr/bin/env python3
"""Idempotent importer for the equipment workbook used by this project."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import psycopg
from openpyxl import load_workbook


SHEETS = {
    'equipment': 'Объекты ремонта',
    'history': 'Ввод начальных данных',
    'dictionaries': 'Справочники',
}

DICT_COLUMNS = {
    'organizations': 'Организации',
    'departments': 'Эксплуатирующее одразделение',
    'work_schedules': 'Графики работы',
    'directions': 'Направления',
    'priorities': 'Приоритет',
    'repair_conditions': 'Условие ремонта',
    'states': 'Текущее состояние',
    'downtime_reasons': 'Причины простоя',
    'repair_works': 'Виды ремонтных работ',
    'repair_work_codes': 'Код вида ремонтных работ',
    'equipment_types': 'Тип оборудования',
    'logic_values': 'Логика',
    'workload_parameters': 'Параметры наработки',
    'units': 'Ед Изм',
    'workload_periods': 'Период плановой наработки',
    'controlled_indicators': 'Контролируемый показатель',
    'repair_contractors': 'Контрагенты исполнители ремонтов',
    'operational_responsibilities': 'Оперативное ведение и управление',
    'warehouses': 'Склады',
}

REQUIRED_HEADERS = {
    'equipment': {
        'Идентификатор для загрузки',
        'Уровень 0',
        'Уровень 1',
        'Уровень 2',
        'Уровень 3',
        'Уровень 4',
        'Уровень 5',
        'Уровень 6',
        'Наименование оборудования',
        'Типовой ОР',
        'Дата ввода в эксплуатацию',
        'Дата выпуска (постройки)',
        'График работы',
        'Срок полезного использования. Месяцев',
        'Инвентарный №',
        'Комментарий',
    },
    'history': {
        'Идентификатор для загрузки',
        'Текущее состояние',
        'Причина простоя',
        'Выполненные ремонтные работы',
        'Дата окончания ремонта',
        'Параметр наработки',
        'Ед Изм',
        'Плановая наработка',
        'Период плановой наработки',
        'Продолжительность периода',
        'Дата работы с',
        'Дата работы по',
        'Значение наработки',
        'Контролируемый показатель',
        'Ед Изм КП',
        'Дата контроля',
        'Значение КП',
    },
    'dictionaries': set(DICT_COLUMNS.values()),
}


def json_default(v: Any):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    return str(v)


def json_dump(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=json_default)


def clean(v: Any):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def load_env_file():
    env_file = Path(__file__).resolve().parents[1] / '.env'
    if not env_file.is_file():
        return

    for line in env_file.read_text(encoding='utf-8').splitlines():
        line = line.strip()
        if not line or line.startswith('#') or '=' not in line:
            continue

        name, value = line.split('=', 1)
        name = name.removeprefix('export ').strip()
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in ('"', "'"):
            value = value[1:-1]
        if name:
            os.environ.setdefault(name, value)


def read_sheet(ws, sheet_key: str):
    values = list(ws.iter_rows(values_only=True))
    if not values:
        raise SystemExit(f'Worksheet {ws.title!r} is empty')

    headers = [
        str(clean(value)) if clean(value) is not None else None
        for value in values[0]
    ]
    named_headers = [header for header in headers if header]
    if len(named_headers) != len(set(named_headers)):
        raise SystemExit(f'Worksheet {ws.title!r} has duplicate column names')

    missing_headers = REQUIRED_HEADERS[sheet_key] - set(named_headers)
    if missing_headers:
        raise SystemExit(
            f'Worksheet {ws.title!r} is missing columns: '
            f'{sorted(missing_headers)}'
        )

    result = []

    for row_num, row in enumerate(values[1:], start=2):
        if not any(clean(x) is not None for x in row):
            continue

        result.append((
            row_num,
            {
                headers[i]: clean(row[i])
                for i in range(min(len(headers), len(row)))
                if headers[i]
            },
        ))

    return result


def text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, Decimal) and v == v.to_integral_value():
        return str(v.quantize(Decimal(1)))
    value = str(v).strip()
    return value or None


def source_text(v):
    return str(v) if v is not None else None


def as_date(v, field: str, row_num: int):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v is None or v == '':
        return None
    if isinstance(v, str):
        try:
            return date.fromisoformat(v.strip())
        except ValueError:
            pass
    raise SystemExit(
        f'Row {row_num}: {field!r} must contain a date, got {v!r}'
    )


def as_number(v, field: str, row_num: int):
    if v is None or v == '':
        return None
    try:
        number = Decimal(str(v))
    except InvalidOperation as error:
        raise SystemExit(
            f'Row {row_num}: {field!r} must contain a number, got {v!r}'
        ) from error
    if not number.is_finite():
        raise SystemExit(
            f'Row {row_num}: {field!r} must contain a finite number, got {v!r}'
        )
    return number


def get_dict_id(cur, group: str, value: Any):
    value = text(value)
    if not value:
        return None

    cur.execute(
        '''
        SELECT id
        FROM dict_import_values
        WHERE dict_group = %s AND value = %s
        ''',
        (group, value),
    )

    row = cur.fetchone()
    return row[0] if row else None


def main():
    load_env_file()
    ap = argparse.ArgumentParser(
        description='Import equipment workbook and preserve relationships'
    )
    default_xlsx = (
        Path(__file__).resolve().parents[1] / 'data' / 'Оборудование.xlsx'
    )
    ap.add_argument('xlsx', type=Path, nargs='?', default=default_xlsx)
    ap.add_argument('--database-url', default=os.getenv('DATABASE_URL'))
    args = ap.parse_args()

    if args.database_url:
        connection_args = {'conninfo': args.database_url}
    else:
        required_env = ('PG_DATABASE', 'PG_USER', 'PG_PASSWORD')
        missing_env = [name for name in required_env if not os.getenv(name)]
        if missing_env:
            raise SystemExit(
                'Set DATABASE_URL or configure '
                + ', '.join(required_env)
                + ' in the environment or backend/.env'
            )
        connection_args = {
            'dbname': os.environ['PG_DATABASE'],
            'user': os.environ['PG_USER'],
            'password': os.environ['PG_PASSWORD'],
            'host': os.getenv('PG_HOST', 'localhost'),
            'port': os.getenv('PG_PORT', '5432'),
        }

    if not args.xlsx.is_file():
        raise SystemExit(f'File not found: {args.xlsx}')

    wb = load_workbook(args.xlsx, data_only=True)

    missing = set(SHEETS.values()) - set(wb.sheetnames)
    if missing:
        raise SystemExit(f'Missing worksheet(s): {sorted(missing)}')

    equipment_rows = read_sheet(wb[SHEETS['equipment']], 'equipment')
    history_rows = read_sheet(wb[SHEETS['history']], 'history')
    dict_rows = read_sheet(wb[SHEETS['dictionaries']], 'dictionaries')

    source_file = args.xlsx.name
    checksum = hashlib.sha256(args.xlsx.read_bytes()).hexdigest()

    # Build the full hierarchy from the seven Level columns.
    # Rows without loading IDs are hierarchy nodes, not disposable rows.
    level_columns = [f'Уровень {i}' for i in range(7)]
    nodes = {}
    equipment_payloads = []

    for row_num, row in equipment_rows:
        parent_key = None

        for depth, column in enumerate(level_columns):
            node_name = text(row.get(column))
            if not node_name:
                continue

            key = (parent_key, node_name)

            if key not in nodes:
                nodes[key] = {
                    'parent_key': parent_key,
                    'name': node_name,
                    'level': depth,
                    'source_row': row_num,
                    'raw': {
                        c: row.get(c)
                        for c in level_columns
                        if row.get(c) is not None
                    },
                }

            parent_key = key

        loading_id = text(row.get('Идентификатор для загрузки'))
        name = text(row.get('Наименование оборудования'))

        if loading_id:
            if not name:
                raise SystemExit(
                    f'Equipment row {row_num} has loading ID '
                    f'{loading_id!r} but no name'
                )

            equipment_payloads.append({
                'row_num': row_num,
                'loading_id': loading_id,
                'loading_id_original': source_text(
                    row.get('Идентификатор для загрузки')
                ),
                'name': name,
                'node_key': parent_key,
                'raw': row,
                'attributes': {
                    k: v
                    for k, v in row.items()
                    if k not in (
                        'Идентификатор для загрузки',
                        'Наименование оборудования',
                    ) and v is not None
                },
            })

    ids = [x['loading_id'] for x in equipment_payloads]

    if len(ids) != len(set(ids)):
        raise SystemExit(
            'Duplicate equipment loading IDs found in workbook; '
            'import stopped to avoid ambiguous links.'
        )

    with psycopg.connect(**connection_args) as conn:
        with conn.cursor() as cur:

            # Import all dictionary categories.
            for group, column in DICT_COLUMNS.items():
                values = {
                    text(row.get(column))
                    for _, row in dict_rows
                    if text(row.get(column))
                }

                for value in values:
                    cur.execute(
                        '''
                        INSERT INTO dict_import_values(dict_group, value)
                        VALUES (%s, %s)
                        ON CONFLICT(dict_group, value) DO NOTHING
                        ''',
                        (group, value),
                    )

            # Populate existing application dictionaries.
            organizations = sorted({
                text(row.get('Организации'))
                for _, row in dict_rows
                if text(row.get('Организации'))
            })

            for value in organizations:
                cur.execute(
                    '''
                    INSERT INTO dict_organizations(name)
                    VALUES (%s)
                    ON CONFLICT(name) DO NOTHING
                    ''',
                    (value,),
                )

            priorities = sorted({
                text(r.get('Приоритет'))
                for _, r in dict_rows
                if text(r.get('Приоритет'))
            })

            for value in priorities:
                cur.execute(
                    '''
                    INSERT INTO dict_priorities(name)
                    VALUES (%s)
                    ON CONFLICT(name) DO NOTHING
                    ''',
                    (value,),
                )

            states = sorted({
                text(r.get('Текущее состояние'))
                for _, r in dict_rows
                if text(r.get('Текущее состояние'))
            })

            for value in states:
                cur.execute(
                    '''
                    INSERT INTO dict_states(name)
                    VALUES (%s)
                    ON CONFLICT(name) DO NOTHING
                    ''',
                    (value,),
                )

            # Only import repair codes explicitly paired with names.
            repair_pairs = set()

            for _, row in dict_rows:
                code = text(row.get('Код вида ремонтных работ'))
                name = text(row.get('Виды ремонтных работ'))

                if code and name:
                    repair_pairs.add((code, name))

            for code, name in sorted(repair_pairs):
                cur.execute(
                    '''
                    INSERT INTO dict_repair_types(code, name)
                    VALUES (%s, %s)
                    ON CONFLICT(code)
                    DO UPDATE SET name = EXCLUDED.name
                    ''',
                    (code, name),
                )

            # Import hierarchy roots -> leaves.
            node_ids = {}
            pending = sorted(
                nodes.items(),
                key=lambda kv: kv[1]['level'],
            )

            for key, node in pending:
                parent_id = (
                    node_ids.get(node['parent_key'])
                    if node['parent_key']
                    else None
                )

                cur.execute(
                    '''
                    SELECT id
                    FROM equipment_hierarchy_nodes
                    WHERE parent_id IS NOT DISTINCT FROM %s
                      AND name = %s
                    LIMIT 1
                    ''',
                    (parent_id, node['name']),
                )

                existing = cur.fetchone()

                if existing:
                    node_id = existing[0]

                    cur.execute(
                        '''
                        UPDATE equipment_hierarchy_nodes
                        SET level = %s,
                            raw_row = %s::jsonb,
                            source_file = %s,
                            source_row_number = %s,
                            updated_at = NOW()
                        WHERE id = %s
                        ''',
                        (
                            node['level'],
                            json_dump(node['raw']),
                            source_file,
                            node['source_row'],
                            node_id,
                        ),
                    )
                else:
                    cur.execute(
                        '''
                        INSERT INTO equipment_hierarchy_nodes(
                            parent_id, name, level, raw_row,
                            source_file, source_row_number
                        )
                        VALUES (%s, %s, %s, %s::jsonb, %s, %s)
                        RETURNING id
                        ''',
                        (
                            parent_id,
                            node['name'],
                            node['level'],
                            json_dump(node['raw']),
                            source_file,
                            node['source_row'],
                        ),
                    )

                    node_id = cur.fetchone()[0]

                node_ids[key] = node_id

            # Import equipment and link it to the hierarchy.
            for item in equipment_payloads:
                row = item['raw']

                schedule_id = get_dict_id(
                    cur,
                    'work_schedules',
                    row.get('График работы'),
                )

                cur.execute(
                    '''
                    INSERT INTO equipment (
                        "loadingId", "loadingIdOriginal", name,
                        attributes, "rawRow", "isContextOnly",
                        hierarchy_node_id, work_schedule_dict_id,
                        typical_repair_code, commissioned_at,
                        manufactured_at, useful_life_months,
                        inventory_number, comment,
                        "createdAt", "updatedAt"
                    )
                    VALUES (
                        %s, %s, %s, %s::jsonb, %s::jsonb, false,
                        %s, %s, %s, %s, %s, %s, %s, %s, NOW(), NOW()
                    )
                    ON CONFLICT ("loadingId") DO UPDATE SET
                        "loadingIdOriginal" = EXCLUDED."loadingIdOriginal",
                        name = EXCLUDED.name,
                        attributes = EXCLUDED.attributes,
                        "rawRow" = EXCLUDED."rawRow",
                        hierarchy_node_id = EXCLUDED.hierarchy_node_id,
                        work_schedule_dict_id = EXCLUDED.work_schedule_dict_id,
                        typical_repair_code = EXCLUDED.typical_repair_code,
                        commissioned_at = EXCLUDED.commissioned_at,
                        manufactured_at = EXCLUDED.manufactured_at,
                        useful_life_months = EXCLUDED.useful_life_months,
                        inventory_number = EXCLUDED.inventory_number,
                        comment = EXCLUDED.comment,
                        "updatedAt" = NOW()
                    RETURNING id
                    ''',
                    (
                        item['loading_id'],
                        item['loading_id_original'],
                        item['name'],
                        json_dump(item['attributes']),
                        json_dump(row),
                        node_ids.get(item['node_key']),
                        schedule_id,
                        text(row.get('Типовой ОР')),
                        as_date(
                            row.get('Дата ввода в эксплуатацию'),
                            'Дата ввода в эксплуатацию',
                            item['row_num'],
                        ),
                        as_date(
                            row.get('Дата выпуска (постройки)'),
                            'Дата выпуска (постройки)',
                            item['row_num'],
                        ),
                        as_number(
                            row.get('Срок полезного использования. Месяцев'),
                            'Срок полезного использования. Месяцев',
                            item['row_num'],
                        ),
                        text(row.get('Инвентарный №')),
                        text(row.get('Комментарий')),
                    ),
                )

                cur.fetchone()

            history_loading_ids = {
                text(row.get('Идентификатор для загрузки'))
                for _, row in history_rows
            }
            relevant_loading_ids = sorted(
                loading_id
                for loading_id in history_loading_ids | set(ids)
                if loading_id
            )
            equipment_ids = {}
            if relevant_loading_ids:
                cur.execute(
                    '''
                    SELECT id, "loadingId"
                    FROM equipment
                    WHERE "loadingId" = ANY(%s)
                    ''',
                    (relevant_loading_ids,),
                )
                equipment_ids = {
                    loading_id: equipment_id
                    for equipment_id, loading_id in cur.fetchall()
                }

            # Import history rows with stable source-row identity.
            matched = unmatched = unmatched_missing_ids = 0

            for row_num, row in history_rows:
                source_id = row.get('Идентификатор для загрузки')
                original_id = text(source_id)
                original_id_text = source_text(source_id)
                equipment_id = equipment_ids.get(original_id)
                is_unmatched = equipment_id is None

                if is_unmatched:
                    unmatched += 1
                    if not original_id:
                        unmatched_missing_ids += 1
                else:
                    matched += 1

                state_name = text(row.get('Текущее состояние'))
                state_id = None

                if state_name:
                    cur.execute(
                        'SELECT id FROM dict_states WHERE name = %s',
                        (state_name,),
                    )
                    found = cur.fetchone()
                    state_id = found[0] if found else None

                repair_name = text(row.get('Выполненные ремонтные работы'))
                repair_type_id = None

                if repair_name:
                    cur.execute(
                        '''
                        SELECT id
                        FROM dict_repair_types
                        WHERE name = %s
                        ORDER BY id
                        LIMIT 1
                        ''',
                        (repair_name,),
                    )
                    found = cur.fetchone()
                    repair_type_id = found[0] if found else None

                cur.execute(
                    '''
                    INSERT INTO equipment_history (
                        "equipmentId", "loadingIdKey", "loadingIdOriginal",
                        raw, "isUnmatched", state_id, downtime_reason,
                        repair_work_name, repair_type_id, repair_completed_at,
                        workload_parameter, workload_unit, planned_workload,
                        planned_workload_period, workload_period_duration,
                        workload_from, workload_to, workload_value,
                        controlled_indicator, controlled_indicator_unit,
                        control_date, controlled_value, source_file,
                        source_sheet, source_row_number,
                        "createdAt", "updatedAt"
                    )
                    VALUES (
                        %s, %s, %s, %s::jsonb, %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s, NOW(), NOW()
                    )
                    ON CONFLICT (
                        source_file, source_sheet, source_row_number
                    ) DO UPDATE SET
                        "equipmentId" = EXCLUDED."equipmentId",
                        "loadingIdKey" = EXCLUDED."loadingIdKey",
                        "loadingIdOriginal" = EXCLUDED."loadingIdOriginal",
                        raw = EXCLUDED.raw,
                        "isUnmatched" = EXCLUDED."isUnmatched",
                        state_id = EXCLUDED.state_id,
                        downtime_reason = EXCLUDED.downtime_reason,
                        repair_work_name = EXCLUDED.repair_work_name,
                        repair_type_id = EXCLUDED.repair_type_id,
                        repair_completed_at = EXCLUDED.repair_completed_at,
                        workload_parameter = EXCLUDED.workload_parameter,
                        workload_unit = EXCLUDED.workload_unit,
                        planned_workload = EXCLUDED.planned_workload,
                        planned_workload_period = EXCLUDED.planned_workload_period,
                        workload_period_duration = EXCLUDED.workload_period_duration,
                        workload_from = EXCLUDED.workload_from,
                        workload_to = EXCLUDED.workload_to,
                        workload_value = EXCLUDED.workload_value,
                        controlled_indicator = EXCLUDED.controlled_indicator,
                        controlled_indicator_unit = EXCLUDED.controlled_indicator_unit,
                        control_date = EXCLUDED.control_date,
                        controlled_value = EXCLUDED.controlled_value,
                        "updatedAt" = NOW()
                    ''',
                    (
                        equipment_id,
                        original_id,
                        original_id_text,
                        json_dump(row),
                        is_unmatched,
                        state_id,
                        text(row.get('Причина простоя')),
                        repair_name,
                        repair_type_id,
                        as_date(
                            row.get('Дата окончания ремонта'),
                            'Дата окончания ремонта',
                            row_num,
                        ),
                        text(row.get('Параметр наработки')),
                        text(row.get('Ед Изм')),
                        as_number(
                            row.get('Плановая наработка'),
                            'Плановая наработка',
                            row_num,
                        ),
                        text(row.get('Период плановой наработки')),
                        text(row.get('Продолжительность периода')),
                        as_date(row.get('Дата работы с'), 'Дата работы с', row_num),
                        as_date(row.get('Дата работы по'), 'Дата работы по', row_num),
                        as_number(
                            row.get('Значение наработки'),
                            'Значение наработки',
                            row_num,
                        ),
                        text(row.get('Контролируемый показатель')),
                        text(row.get('Ед Изм КП')),
                        as_date(row.get('Дата контроля'), 'Дата контроля', row_num),
                        as_number(row.get('Значение КП'), 'Значение КП', row_num),
                        source_file,
                        SHEETS['history'],
                        row_num,
                    ),
                )

            # One import report per file checksum.
            report = {
                'source_file': source_file,
                'sha256': checksum,
                'equipment_rows': len(equipment_payloads),
                'hierarchy_nodes': len(nodes),
                'history_rows': len(history_rows),
                'matched_history_rows': matched,
                'unmatched_history_rows': unmatched,
                'unmatched_history_rows_without_id': unmatched_missing_ids,
                'dictionary_categories': len(DICT_COLUMNS),
                'notes': [
                    'Requests, telemetry and incident episodes are not fabricated: '
                    'the workbook has no corresponding source records.'
                ],
            }

            cur.execute(
                '''
                INSERT INTO equipment_import_reports (
                    "unmatchedCount", "unmatchedMissingIdsCount",
                    "matchedCount", report, source_file, source_sha256,
                    "createdAt", "updatedAt"
                )
                VALUES (%s, %s, %s, %s::jsonb, %s, %s, NOW(), NOW())
                ON CONFLICT (source_sha256) DO UPDATE SET
                    "unmatchedCount" = EXCLUDED."unmatchedCount",
                    "unmatchedMissingIdsCount" = EXCLUDED."unmatchedMissingIdsCount",
                    "matchedCount" = EXCLUDED."matchedCount",
                    report = EXCLUDED.report,
                    source_file = EXCLUDED.source_file,
                    "updatedAt" = NOW()
                ''',
                (
                    unmatched,
                    unmatched_missing_ids,
                    matched,
                    json_dump(report),
                    source_file,
                    checksum,
                ),
            )

    print(f'Equipment rows imported: {len(equipment_payloads)}')
    print(f'Hierarchy nodes imported: {len(nodes)}')
    print(
        f'History rows: {len(history_rows)}; '
        f'matched: {matched}; unmatched: {unmatched}'
    )
    print(
        f'All {len(DICT_COLUMNS)} dictionary categories imported '
        'to dict_import_values.'
    )


if __name__ == '__main__':
    main()